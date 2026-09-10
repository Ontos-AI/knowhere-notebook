#!/usr/bin/env python3
"""Stratified pilot sample from 观心 knowhere自测集.xlsx.

Takes the first PILOT_CASES_PER_STRATUM rows (by seq_id) from each
(sheet, 难度) bucket. Does not call Knowhere or write the case workspace.
"""
from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX = Path(
    "/Users/wuchengke/Desktop/skills-coding/观心2.0-RAG-v1.1-demo/knowhere自测集.xlsx"
)
DEFAULT_OUT = ROOT / ".tmp" / "guanxin-pilot-cases.json"
PILOT_CASES_PER_STRATUM = 4


def load_rows(xlsx: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows: list[dict] = []
    try:
        for sheet_name in wb.sheetnames:
            sheet_rows = list(wb[sheet_name].iter_rows(values_only=True))
            if not sheet_rows:
                continue
            header = [
                str(h).strip() if h is not None else f"col{i}"
                for i, h in enumerate(sheet_rows[0])
            ]
            for raw in sheet_rows[1:]:
                if not raw or raw[0] in (None, ""):
                    continue
                item = {
                    header[i]: (raw[i] if i < len(raw) else None)
                    for i in range(len(header))
                }
                query = str(item.get("具体query") or "").strip()
                if not query:
                    continue
                rows.append(
                    {
                        "sheet": sheet_name,
                        "seq_id": str(item.get("seq_id") or ""),
                        "query": query,
                        "disease": str(item.get("具体疾病名称") or "").strip(),
                        "scene": str(item.get("应用场景-考察能力") or "").strip(),
                        "difficulty": str(item.get("难度") or "").strip(),
                        "input_type": str(item.get("输入类型") or "").strip(),
                    }
                )
    finally:
        wb.close()
    return rows


def sample(rows: list[dict]) -> list[dict]:
    buckets: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in rows:
        buckets[(row["sheet"], row["difficulty"] or "?")].append(row)
    picked: list[dict] = []
    for key in sorted(buckets):
        group = sorted(buckets[key], key=lambda row: row["seq_id"])
        picked.extend(group[:PILOT_CASES_PER_STRATUM])
    return sorted(picked, key=lambda row: (row["sheet"], row["seq_id"]))


def main() -> None:
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    rows = load_rows(xlsx)
    picked = sample(rows)
    out.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "xlsx": str(xlsx),
        "per_stratum": PILOT_CASES_PER_STRATUM,
        "source_count": len(rows),
        "pilot_count": len(picked),
        "cases": picked,
    }
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {len(picked)} / {len(rows)} -> {out}")


if __name__ == "__main__":
    main()
